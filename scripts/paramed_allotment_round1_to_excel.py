#!/usr/bin/env python3
"""
Paramed Allotment Round 1 PDF → Excel Converter
Reads database/paramed_allotment_round1.pdf and writes results/paramed_allotment_round1.xlsx

Column order: SNO, GRANK, APP. NO, NAME, COMMUNITY, T.MARK, R1 ALLOTED, R2 ALLOTED, CATEGORY(PWD)
"""

import os
import sys
import pdfplumber
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

# ── Paths ────────────────────────────────────────────────────────────────────
BASE_DIR   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PDF_PATH   = os.path.join(BASE_DIR, "database", "paramed_allotment_round1.pdf")
RESULT_DIR = os.path.join(BASE_DIR, "results")
OUT_PATH   = os.path.join(RESULT_DIR, "paramed_allotment_round1.xlsx")

os.makedirs(RESULT_DIR, exist_ok=True)

COLUMNS = [
    "SNO",
    "GRANK",
    "APP. NO",
    "NAME",
    "COMMUNITY",
    "T.MARK",
    "R1 ALLOTED",
    "R2 ALLOTED",
    "CATEGORY(PWD)",
]

# ── Style helpers ─────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="1A3A5C")   # dark navy
HEADER_FONT   = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
ODD_ROW_FILL  = PatternFill("solid", fgColor="EEF4FB")   # light blue-grey
EVEN_ROW_FILL = PatternFill("solid", fgColor="FFFFFF")
DATA_FONT     = Font(name="Calibri", size=10)
TITLE_FONT    = Font(bold=True, name="Calibri", size=13, color="1A3A5C")
THIN_BORDER   = Border(
    left   = Side(style="thin", color="CCCCCC"),
    right  = Side(style="thin", color="CCCCCC"),
    top    = Side(style="thin", color="CCCCCC"),
    bottom = Side(style="thin", color="CCCCCC"),
)

COL_WIDTHS = {
    "SNO":           8,
    "GRANK":         10,
    "APP. NO":       16,
    "NAME":          30,
    "COMMUNITY":     14,
    "T.MARK":        12,
    "R1 ALLOTED":    46,
    "R2 ALLOTED":    46,
    "CATEGORY(PWD)": 18,
}


def clean_str(val) -> str:
    """Trim and collapse internal whitespace for single-line fields."""
    if val is None:
        return ""
    return " ".join(str(val).split())


def clean_multiline(val) -> str:
    """Trim individual lines while preserving multiline structure for college/course."""
    if val is None:
        return ""
    lines = [" ".join(line.split()) for line in str(val).splitlines()]
    return "\n".join(l for l in lines if l)


# ── Extract rows from PDF ─────────────────────────────────────────────────────
def extract_rows(pdf_path: str) -> list[list]:
    rows = []

    with pdfplumber.open(pdf_path) as pdf:
        total = len(pdf.pages)
        print(f"  PDF has {total} pages. Extracting…")

        for page_num, page in enumerate(pdf.pages, start=1):
            if page_num % 50 == 0 or page_num == 1:
                print(f"  → Page {page_num}/{total}", end="\r", flush=True)

            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    if not row or not any(row):
                        continue

                    first_col = clean_str(row[0]).upper()
                    if first_col in ("SNO", "SNO GRANK", "") or "SNO" in first_col:
                        continue

                    # Round 1 tables have 8 columns:
                    # [0: SNO, 1: GRANK, 2: APP. NO, 3: NAME, 4: COMMUNITY, 5: T.MARK, 6: ALLOTTED TO, 7: CATEGORY]
                    if len(row) < 7:
                        continue

                    sno = clean_str(row[0])
                    # SNO should be numeric
                    if not sno.isdigit():
                        continue

                    grank = clean_str(row[1])
                    app_no = clean_str(row[2])
                    name = clean_str(row[3])
                    community = clean_str(row[4])
                    tmark = clean_str(row[5])
                    r1_allotted = clean_multiline(row[6])
                    r2_allotted = ""  # Not applicable in Round 1
                    category = clean_str(row[7]) if len(row) > 7 else ""

                    rows.append([
                        sno,
                        grank,
                        app_no,
                        name,
                        community,
                        tmark,
                        r1_allotted,
                        r2_allotted,
                        category,
                    ])

    print(f"\n  ✔ Extracted {len(rows):,} data rows")
    return rows


# ── Write Excel ───────────────────────────────────────────────────────────────
def write_excel(rows: list[list], out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Round 1 Allotment"

    # Title row
    title_text = "Paramedical Degree Courses 2026-2027 — Provisional List of Candidates Allotted in Round 1"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font      = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_cell.fill      = PatternFill("solid", fgColor="D6E4F5")
    ws.row_dimensions[1].height = 36

    # Header row
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=col_name)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = THIN_BORDER
    ws.row_dimensions[2].height = 24

    # Freeze pane below header
    ws.freeze_panes = "A3"

    # Data rows
    for r_idx, row in enumerate(rows, start=3):
        fill = ODD_ROW_FILL if (r_idx % 2 == 1) else EVEN_ROW_FILL
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font      = DATA_FONT
            cell.fill      = fill
            cell.border    = THIN_BORDER

            col_name = COLUMNS[c_idx - 1]
            if col_name in ("SNO", "GRANK", "APP. NO", "COMMUNITY", "T.MARK", "CATEGORY(PWD)"):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_name in ("R1 ALLOTED", "R2 ALLOTED"):
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

            # Numeric coercion
            if col_name in ("SNO", "GRANK") and value.isdigit():
                cell.value = int(value)
            elif col_name == "T.MARK":
                try:
                    cell.value = float(value) if "." in value else int(value)
                except ValueError:
                    pass

    # Column widths
    for c_idx, col_name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = COL_WIDTHS[col_name]

    # Auto-filter on header
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}2"

    wb.save(out_path)
    print(f"  ✔ Saved → {out_path}")


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    if not os.path.exists(PDF_PATH):
        print(f"ERROR: PDF not found at {PDF_PATH}")
        sys.exit(1)

    print(f"\n📄 Source : {PDF_PATH}")
    print(f"📊 Output : {OUT_PATH}\n")

    rows = extract_rows(PDF_PATH)

    if not rows:
        print("ERROR: No data rows extracted. Check PDF structure.")
        sys.exit(1)

    print("\nWriting Excel…")
    write_excel(rows, OUT_PATH)
    print(f"\n✅ Done! {len(rows):,} records written to:\n   {OUT_PATH}\n")


if __name__ == "__main__":
    main()
