#!/usr/bin/env python3
"""
DGNM Science Ranklist PDF → Excel Converter
Reads database/dgnm_science_ranklist.pdf and writes results/dgnm_science_ranklist.xlsx
Columns: RANK, APP.NO, NAME, COMMUNITY, TOTAL MARKS, COMMUNITY RANK
"""

import os
import sys
import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Paths ─────────────────────────────────────────────────────────────────────
BASE_DIR   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PDF_PATH   = os.path.join(BASE_DIR, "database", "dgnm_science_ranklist.pdf")
RESULT_DIR = os.path.join(BASE_DIR, "results")
OUT_PATH   = os.path.join(RESULT_DIR, "dgnm_science_ranklist.xlsx")

os.makedirs(RESULT_DIR, exist_ok=True)

COLUMNS = ["RANK", "APP.NO", "NAME", "COMMUNITY", "TOTAL MARKS", "COMMUNITY RANK"]

# ── Styles ─────────────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="1A3A5C")   # dark navy
HEADER_FONT   = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
ODD_ROW_FILL  = PatternFill("solid", fgColor="EEF4FB")   # light blue-grey
EVEN_ROW_FILL = PatternFill("solid", fgColor="FFFFFF")
DATA_FONT     = Font(name="Calibri", size=10)
TITLE_FONT    = Font(bold=True, name="Calibri", size=13, color="1A3A5C")
THIN_BORDER   = Border(
    left   = Side(style="thin", color="CCCCCC"),
    right  = Side(style="thin", color="CCCCCC"),
    top    = Side(style="thin", color="CCCCCC"),
    bottom = Side(style="thin", color="CCCCCC"),
)
COL_WIDTHS = {
    "RANK":           8,
    "APP.NO":         16,
    "NAME":           32,
    "COMMUNITY":      14,
    "TOTAL MARKS":    14,
    "COMMUNITY RANK": 16,
}

# ── Extract rows ──────────────────────────────────────────────────────────────
def extract_rows(pdf_path: str) -> list[list]:
    rows = []
    with pdfplumber.open(pdf_path) as pdf:
        total = len(pdf.pages)
        print(f"  PDF has {total} pages. Extracting…")
        for page_num, page in enumerate(pdf.pages, start=1):
            if page_num % 50 == 0 or page_num == 1:
                print(f"  → Page {page_num}/{total}", end="\r", flush=True)
            for table in page.extract_tables():
                for row in table:
                    clean = [
                        " ".join(str(c).split()) if c is not None else ""
                        for c in row
                    ]
                    # Skip header / empty rows
                    if clean[0].upper() in ("RANK", ""):
                        continue
                    if len(clean) != len(COLUMNS):
                        continue
                    if not clean[0].isdigit():
                        continue
                    rows.append(clean)
    print(f"\n  ✔ Extracted {len(rows):,} data rows")
    return rows

# ── Write Excel ───────────────────────────────────────────────────────────────
def write_excel(rows: list[list], out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DGNM Science Ranklist"

    # Title row
    title_text = "Diploma in General Nursing & Midwifery 2026-2027 — Provisional Rank List (Science Group)"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    tc = ws.cell(row=1, column=1, value=title_text)
    tc.font      = TITLE_FONT
    tc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    tc.fill      = PatternFill("solid", fgColor="D6E4F5")
    ws.row_dimensions[1].height = 36

    # Header row
    for ci, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=ci, value=col)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = THIN_BORDER
    ws.row_dimensions[2].height = 22
    ws.freeze_panes = "A3"

    # Data rows
    for ri, row in enumerate(rows, start=3):
        fill = ODD_ROW_FILL if (ri % 2 == 1) else EVEN_ROW_FILL
        for ci, value in enumerate(row, start=1):
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.font   = DATA_FONT
            cell.fill   = fill
            cell.border = THIN_BORDER
            col_name = COLUMNS[ci - 1]
            if col_name in ("RANK", "TOTAL MARKS", "COMMUNITY RANK", "APP.NO"):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            # Numeric coercion
            if col_name == "RANK" and value.isdigit():
                cell.value = int(value)
            elif col_name == "TOTAL MARKS":
                try:
                    cell.value = float(value)
                except ValueError:
                    pass
            elif col_name == "COMMUNITY RANK" and value.isdigit():
                cell.value = int(value)

    # Column widths & auto-filter
    for ci, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTHS[col]
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}2"

    wb.save(out_path)
    print(f"  ✔ Saved → {out_path}")

# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    if not os.path.exists(PDF_PATH):
        print(f"ERROR: PDF not found at {PDF_PATH}")
        sys.exit(1)

    print(f"\n📄 Source : {PDF_PATH}")
    print(f"📊 Output : {OUT_PATH}\n")

    rows = extract_rows(PDF_PATH)
    if not rows:
        print("ERROR: No data rows extracted.")
        sys.exit(1)

    print("\nWriting Excel…")
    write_excel(rows, OUT_PATH)
    print(f"\n✅ Done! {len(rows):,} records written to:\n   {OUT_PATH}\n")

if __name__ == "__main__":
    main()
